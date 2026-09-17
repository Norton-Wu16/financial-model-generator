"""Excel 单元格样式定义（财务建模惯例）。

约定:
- 输入格: 黄色填充 + 蓝色字体（用户可编辑）
- 公式格: 黑色字体（自动计算）
- 表头: 深色填充 + 白色加粗字体
- 标签列: 浅灰填充
- 链接格: 黑色字体（跨表引用，可斜体标识）
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter

# 填充
INPUT_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")      # 浅黄
INPUT_FILL_STRONG = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")  # 亮黄
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")       # 深蓝
SUBHEADER_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")    # 浅蓝
LABEL_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")        # 浅灰
TOTAL_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")       # 浅绿（合计）
CHECK_OK_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")     # 绿
CHECK_BAD_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")   # 红

# 字体
INPUT_FONT = Font(name="Calibri", size=10, color="0000FF")        # 蓝色（输入）
FORMULA_FONT = Font(name="Calibri", size=10, color="000000")      # 黑色（公式）
LINK_FONT = Font(name="Calibri", size=10, color="000000", italic=True)  # 斜体（跨表链接）
HEADER_FONT = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
SUBHEADER_FONT = Font(name="Calibri", size=10, color="1F4E78", bold=True)
LABEL_FONT = Font(name="Calibri", size=10, color="000000")
TOTAL_FONT = Font(name="Calibri", size=10, color="000000", bold=True)
TITLE_FONT = Font(name="Calibri", size=14, color="1F4E78", bold=True)

# 对齐
LEFT_ALIGN = Alignment(horizontal="left", vertical="center")
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
RIGHT_ALIGN = Alignment(horizontal="right", vertical="center")
WRAP_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)

# 边框
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
THICK_BOTTOM = Border(left=THIN, right=THIN, top=THIN,
                       bottom=Side(border_style="medium", color="000000"))

# 数字格式
NUM_FMT = '#,##0.00;[Red](#,##0.00);"-"'              # 两位小数带千分位
NUM_FMT_INT = '#,##0;[Red](#,##0);"-"'                 # 整数
PCT_FMT = '0.0%;[Red](0.0%);"-"'                       # 百分比一位小数
MULT_FMT = '0.00"x"'                                  # 倍数
PRICE_FMT = '#,##0.00'                                # 价格


def style_input(cell, fmt=NUM_FMT):
    """标记为输入单元格。"""
    cell.fill = INPUT_FILL
    cell.font = INPUT_FONT
    cell.number_format = fmt
    cell.alignment = RIGHT_ALIGN
    cell.border = BORDER


def style_formula(cell, fmt=NUM_FMT, link=False):
    """标记为公式单元格。link=True 表示跨表引用。"""
    cell.font = LINK_FONT if link else FORMULA_FONT
    cell.number_format = fmt
    cell.alignment = RIGHT_ALIGN
    cell.border = BORDER


def style_total(cell, fmt=NUM_FMT):
    """标记为合计行。"""
    cell.font = TOTAL_FONT
    cell.fill = TOTAL_FILL
    cell.number_format = fmt
    cell.alignment = RIGHT_ALIGN
    cell.border = BORDER


def style_header(cell):
    """表头。"""
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = CENTER_ALIGN
    cell.border = BORDER


def style_subheader(cell):
    """次级表头。"""
    cell.fill = SUBHEADER_FILL
    cell.font = SUBHEADER_FONT
    cell.alignment = LEFT_ALIGN
    cell.border = BORDER


def style_label(cell, indent=0):
    """行标签。"""
    cell.fill = LABEL_FILL
    cell.font = LABEL_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=indent)
    cell.border = BORDER


def style_title(cell):
    cell.font = TITLE_FONT
    cell.alignment = LEFT_ALIGN


def write_header_row(ws, row, headers, start_col=1):
    """写一行表头。"""
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        style_header(c)


def set_col_widths(ws, widths):
    """widths: dict {col_letter: width}"""
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
